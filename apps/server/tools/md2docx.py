"""Xuất bài viết Markdown ra Word (.docx) cho VsScience, kèm trích dẫn và danh mục tài liệu tham khảo.

  python md2docx.py <bai-viet.md> [--out <bai-viet.docx>] [--bib <thu-vien.bib|.ris>] [--style apa|vancouver]

Markdown hỗ trợ: tiêu đề (#…######), đoạn văn, **đậm**, *nghiêng*, `mã`, [link](url), danh sách (-, *, 1.),
trích dẫn khối (>), bảng (| a | b |), ảnh trên một dòng riêng ![chú thích](duong/dan.png), khối mã ```.
Trích dẫn: [@khoa] hoặc [@khoa1; @khoa2, tr. 12] → APA: (Tác giả, năm) — Vancouver: [1], [1,2].
Cuối bài tự thêm mục "Tài liệu tham khảo" gồm các mục đã trích (nếu có --bib).
"""

import argparse
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import vs_refs  # noqa: E402

try:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor
except ImportError:
    print("LỖI: Máy chưa có thư viện python-docx.", file=sys.stderr)
    sys.exit(1)

CITE = re.compile(r"\[(@[^\]]+)\]")
INLINE = re.compile(r"(\*\*[^*]+\*\*|__[^_]+__|\*[^*\s][^*]*\*|_[^_\s][^_]*_|`[^`]+`|\[[^\]]+\]\([^)]+\))")


class Citations:
    def __init__(self, entries: list[dict], style: str):
        self.by_key = {e["key"]: e for e in entries}
        self.style = style
        self.order: list[str] = []
        self.missing: list[str] = []

    def replace(self, text: str) -> str:
        def sub(m):
            items = [x.strip() for x in m.group(1).split(";") if x.strip()]
            keys, locs = [], []
            for it in items:
                km = re.match(r"@([\w:.\-/]+)\s*,?\s*(.*)$", it)
                if not km:
                    continue
                k, loc = km.group(1), km.group(2).strip()
                if k not in self.by_key:
                    self.missing.append(k)
                    return m.group(0)
                if k not in self.order:
                    self.order.append(k)
                keys.append(k)
                locs.append(loc)
            if not keys:
                return m.group(0)
            if self.style == "vancouver":
                nums = [str(self.order.index(k) + 1) for k in keys]
                loc = next((l for l in locs if l), "")
                return f"[{','.join(nums)}{', ' + loc if loc else ''}]"
            return vs_refs.intext_apa([self.by_key[k] for k in keys], locs)

        return CITE.sub(sub, text) if self.by_key else text


def add_inline(par, text: str) -> None:
    for part in INLINE.split(text):
        if not part:
            continue
        if (part.startswith("**") and part.endswith("**")) or (part.startswith("__") and part.endswith("__")):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
        elif re.fullmatch(r"\[[^\]]+\]\([^)]+\)", part):
            label = re.match(r"\[([^\]]+)\]", part).group(1)
            r = par.add_run(label)
            r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x99)
            r.underline = True
        elif len(part) > 2 and part[0] in "*_" and part[-1] == part[0]:
            par.add_run(part[1:-1]).italic = True
        else:
            par.add_run(part)


def table_rows(lines: list[str]) -> list[list[str]]:
    rows = []
    for ln in lines:
        cells = [c.strip() for c in ln.strip().strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
            continue
        rows.append(cells)
    return rows


def convert(md_path: str, out: str, bib: str | None, style: str) -> None:
    with open(md_path, encoding="utf-8-sig") as f:
        lines = f.read().splitlines()
    cites = Citations(vs_refs.load(bib) if bib else [], style)
    base = os.path.dirname(os.path.abspath(md_path))
    doc = docx.Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(13)
    i = 0
    has_title = False
    while i < len(lines):
        ln = lines[i]
        s = ln.strip()
        if not s:
            i += 1
            continue
        if s.startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            r = p.add_run("\n".join(block))
            r.font.name = "Consolas"
            r.font.size = Pt(10)
            continue
        h = re.match(r"^(#{1,6})\s+(.*)$", s)
        if h:
            n = len(h.group(1))
            text = cites.replace(h.group(2).strip())
            if n == 1 and not has_title and not doc.paragraphs:
                # "# Tên bài" ở đầu file là tiêu đề bài; "##" khi đó là mục cấp 1.
                doc.add_heading(text, level=0)
                has_title = True
            else:
                doc.add_heading(text, level=max(1, min(4, n - 1 if has_title else n)))
            i += 1
            continue
        img = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", s)
        if img:
            src = img.group(2).strip().strip("<>")
            full = src if os.path.isabs(src) else os.path.join(base, src)
            if os.path.isfile(full):
                doc.add_picture(full, width=Cm(15))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                if img.group(1):
                    cap = doc.add_paragraph()
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap.add_run(cites.replace(img.group(1))).italic = True
            else:
                doc.add_paragraph(f"[Không thấy ảnh: {src}]")
            i += 1
            continue
        if s.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            rows = table_rows(block)
            if rows:
                ncol = max(len(r) for r in rows)
                t = doc.add_table(rows=len(rows), cols=ncol)
                t.style = "Table Grid"
                for ri, row in enumerate(rows):
                    for ci in range(ncol):
                        cell = t.cell(ri, ci)
                        cell.text = ""
                        add_inline(cell.paragraphs[0], cites.replace(row[ci] if ci < len(row) else ""))
                        if ri == 0:
                            for r in cell.paragraphs[0].runs:
                                r.bold = True
            continue
        lm = re.match(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$", ln)
        if lm:
            style_name = "List Number" if lm.group(2)[0].isdigit() else "List Bullet"
            level = len(lm.group(1).replace("\t", "    ")) // 2
            if level:
                style_name += f" {min(level + 1, 3)}"
            try:
                p = doc.add_paragraph(style=style_name)
            except KeyError:
                p = doc.add_paragraph(style="List Number" if lm.group(2)[0].isdigit() else "List Bullet")
            add_inline(p, cites.replace(lm.group(3)))
            i += 1
            continue
        if s.startswith(">"):
            quote = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote.append(lines[i].strip()[1:].strip())
                i += 1
            p = doc.add_paragraph(style="Quote") if "Quote" in [st.name for st in doc.styles] else doc.add_paragraph()
            add_inline(p, cites.replace(" ".join(quote)))
            continue
        para = [s]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#{1,6}\s|[-*+]\s|\d+[.)]\s|\||>|```|!\[)", lines[i].strip()):
            para.append(lines[i].strip())
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline(p, cites.replace(" ".join(para)))

    if cites.order:
        doc.add_heading("Tài liệu tham khảo", level=1)
        chosen = [cites.by_key[k] for k in cites.order]
        if style == "apa":
            chosen = vs_refs.sort_apa(chosen)
        for n, e in enumerate(chosen, 1):
            p = doc.add_paragraph()
            if style == "apa":
                p.paragraph_format.left_indent = Cm(1.27)
                p.paragraph_format.first_line_indent = Cm(-1.27)
            else:
                p.add_run(f"{n}. ")
            for text, italic in vs_refs.format_entry(e, style):
                r = p.add_run(text)
                r.italic = italic
    doc.save(out)
    print(f"Đã xuất → {out}")
    if cites.order:
        print(f"Đã trích {len(cites.order)} tài liệu ({style.upper()}).")
    for k in dict.fromkeys(cites.missing):
        print(f"- Không có khóa '{k}' trong thư viện: giữ nguyên [@{k}] trong bài, cần thêm vào thư viện.")


def main() -> None:
    ap = argparse.ArgumentParser(description="Markdown → Word")
    ap.add_argument("markdown")
    ap.add_argument("--out")
    ap.add_argument("--bib")
    ap.add_argument("--style", choices=vs_refs.STYLES, default="apa")
    args = ap.parse_args()
    if not os.path.isfile(args.markdown):
        print(f"LỖI: Không thấy file: {args.markdown}", file=sys.stderr)
        sys.exit(1)
    if args.bib and not os.path.isfile(args.bib):
        print(f"LỖI: Không thấy thư viện tài liệu: {args.bib}", file=sys.stderr)
        sys.exit(1)
    out = args.out or os.path.splitext(args.markdown)[0] + ".docx"
    convert(args.markdown, out, args.bib, args.style)


if __name__ == "__main__":
    main()
