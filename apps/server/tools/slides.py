"""Đọc PowerPoint (.pptx) cho Claude trong VsScience (dùng python-pptx đi kèm app).

  python slides.py info   <file.pptx>                       số slide, tiêu đề từng slide
  python slides.py text   <file.pptx> [--pages 3-7,10]      chữ, bảng, ghi chú người thuyết trình của từng slide
  python slides.py search <file.pptx> "<từ khóa>"           các slide có từ khóa, kèm đoạn trích

Số slide là số thứ tự trong bài (slide 1 là slide đầu tiên). Slide ẩn có ghi "(ẩn)".
"""

import argparse
import os
import sys
import unicodedata

MAX_TEXT = 60_000


def fail(msg: str) -> None:
    print(f"LỖI: {msg}", file=sys.stderr)
    sys.exit(1)


try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError:
    fail("Máy chưa có thư viện python-pptx để đọc PowerPoint.")


def nfc(s: str) -> str:
    return unicodedata.normalize("NFC", s or "").strip()


def open_pptx(path: str):
    if not os.path.isfile(path):
        fail(f"Không thấy file: {path}")
    if path.lower().endswith(".ppt"):
        fail("File PowerPoint đời cũ (.ppt): mở bằng PowerPoint và lưu lại thành .pptx rồi đọc lại.")
    try:
        return Presentation(path)
    except Exception as e:  # noqa: BLE001
        fail(f"Không mở được file (hỏng hoặc không phải .pptx): {e}")


def parse_pages(spec: str | None, count: int) -> list[int]:
    if not spec:
        return list(range(1, count + 1))
    pages: list[int] = []
    for part in spec.replace(" ", "").split(","):
        if not part:
            continue
        if "-" in part:
            a, b = part.split("-", 1)
            pages.extend(range(int(a) if a else 1, (int(b) if b else count) + 1))
        else:
            pages.append(int(part))
    seen = set()
    return [p for p in pages if 1 <= p <= count and not (p in seen or seen.add(p))]


def title_of(slide) -> str:
    try:
        if slide.shapes.title is not None:
            return nfc(slide.shapes.title.text_frame.text)
    except Exception:  # noqa: BLE001
        pass
    return ""


def shape_lines(shape, depth: int = 0) -> list[str]:
    """Chữ trong một hình: khung chữ (giữ cấp gạch đầu dòng), bảng (mỗi hàng một dòng), nhóm hình, chú thích ảnh."""
    lines: list[str] = []
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for s in shape.shapes:
            lines.extend(shape_lines(s, depth))
        return lines
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [nfc(c.text).replace("\n", " ") for c in row.cells]
            lines.append("| " + " | ".join(cells) + " |")
        return lines
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            t = nfc(p.text.replace("\v", " "))
            if t:
                lines.append("  " * (p.level or 0) + ("- " if p.level else "") + t)
        return lines
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        desc = ""
        try:
            desc = nfc(shape._element.nvPicPr.cNvPr.get("descr") or "")
        except Exception:  # noqa: BLE001
            pass
        lines.append(f"[Hình{': ' + desc if desc else ''}]")
    if getattr(shape, "has_chart", False) and shape.has_chart:
        try:
            ch = shape.chart
            name = nfc(ch.chart_title.text_frame.text) if ch.has_title else ""
        except Exception:  # noqa: BLE001
            name = ""
        lines.append(f"[Biểu đồ{': ' + name if name else ''}]")
    return lines


def slide_block(n: int, slide) -> str:
    hidden = slide._element.get("show") == "0"
    title = title_of(slide)
    out = [f"=== Slide {n}{' (ẩn)' if hidden else ''}{': ' + title if title else ''} ==="]
    for shape in slide.shapes:
        if slide.shapes.title is not None and shape.shape_id == slide.shapes.title.shape_id:
            continue
        out.extend(shape_lines(shape))
    if slide.has_notes_slide:
        notes = nfc(slide.notes_slide.notes_text_frame.text if slide.notes_slide.notes_text_frame else "")
        if notes:
            out.append(f"Ghi chú người thuyết trình: {notes}")
    return "\n".join(out)


def cmd_info(prs, _args) -> None:
    slides = list(prs.slides)
    print(f"Số slide: {len(slides)}")
    for i, s in enumerate(slides, 1):
        hidden = " (ẩn)" if s._element.get("show") == "0" else ""
        print(f"- Slide {i}{hidden}: {title_of(s) or '(không có tiêu đề)'}")


def cmd_text(prs, args) -> None:
    slides = list(prs.slides)
    pages = parse_pages(args.pages, len(slides))
    out, used = [], 0
    for n in pages:
        block = slide_block(n, slides[n - 1]) + "\n"
        if used + len(block) > MAX_TEXT and out:
            out.append(f"[Dừng ở trước slide {n} vì đã dài; đọc tiếp bằng --pages {n}-{pages[-1]}]")
            break
        out.append(block)
        used += len(block)
    print("\n".join(out) if out else "Không có slide nào trong phạm vi yêu cầu.")


def cmd_search(prs, args) -> None:
    query = unicodedata.normalize("NFC", args.query).casefold()
    hits = 0
    for n, slide in enumerate(prs.slides, 1):
        text = slide_block(n, slide)
        low = text.casefold()
        i = low.find(query)
        if i < 0:
            continue
        hits += 1
        a, b = max(0, i - 100), min(len(text), i + len(query) + 100)
        print(f"- Slide {n} ({low.count(query)} lần): …{' '.join(text[a:b].split())}…")
    if not hits:
        print(f"Không thấy “{args.query}” trong bài trình chiếu.")


def main() -> None:
    parser = argparse.ArgumentParser(description="Đọc PowerPoint cho Claude")
    sub = parser.add_subparsers(dest="cmd", required=True)
    for name in ("info", "text", "search"):
        p = sub.add_parser(name)
        p.add_argument("file")
        if name == "search":
            p.add_argument("query")
        if name == "text":
            p.add_argument("--pages")
    args = parser.parse_args()
    prs = open_pptx(args.file)
    {"info": cmd_info, "text": cmd_text, "search": cmd_search}[args.cmd](prs, args)


if __name__ == "__main__":
    try:
        main()
    except ValueError as e:
        fail(f"Tham số không hợp lệ: {e}")
